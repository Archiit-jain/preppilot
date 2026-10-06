"""
Build the PrepPilot quick report: a short summary of the project, with screenshots.

    python docs/dpr/build_quick.py

Writes docs/PrepPilot_Quick_Report.docx and docs/PrepPilot_Quick_Report.pdf. It
reuses the page layout and helpers of build_dpr.py, and draws its facts from the
same sources as the full report (docs/PrepPilot_DPR.pdf).
"""

import os
import shutil
import subprocess
import sys

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_dpr import (  # noqa: E402
    ACCENT, ASSETS, AUTHOR, DATE, DOCS, GREY, LAND_W, REPO, SUBTITLE, TITLE, WEB,
    Report, diagram, shot, to_pdf,
)

OUT_DOCX = os.path.join(DOCS, "PrepPilot_Quick_Report.docx")
OUT_PDF = os.path.join(DOCS, "PrepPilot_Quick_Report.pdf")


def build():
    R = Report({}, header="PrepPilot — Quick Report")
    d = R.doc
    # The cover is part of page 1 here, so page 1 keeps its header and page number.
    d.sections[0].different_first_page_header_footer = False

    # ---------------------------------------------------------- title block
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(ASSETS, "branding", "preppilot-icon.png"), width=Inches(0.8))
    for text, size, bold, color in (
        (TITLE, 26, True, ACCENT),
        (SUBTITLE, 13, False, None),
        (f"Quick Report  ·  {AUTHOR}  ·  {DATE}", 11, False, GREY),
    ):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color

    # ---------------------------------------------------------- 1
    R.h1("What PrepPilot Is")
    R.para(
        "PrepPilot helps a student organise an examination syllabus as Subject → Chapter → "
        "Topic → optional Sub-topic, tick off what is genuinely finished, time study sessions, "
        "revise, and see honest progress — with or without an internet connection. Timers, "
        "flashcards, notes, quizzes, reminders and the AI assistant all attach to that same "
        "syllabus, so effort, revision and progress stay connected instead of being spread over "
        "several apps. Its principle is **Tracker first, AI second**: the tracker is complete "
        "without the network, and the assistant can never change the syllabus without the "
        "student's confirmation."
    )
    R.table("Project at a glance", [
        ["Item", "Details"],
        ["What it is", "Offline-first study tracker with optional AI help"],
        ["Platform", "Android app; the same code also runs as a browser demo (preppilottdemo.netlify.app)"],
        ["Built with", "React Native, Expo and TypeScript; SQLite with Drizzle ORM on the device"],
        ["Backend", "Supabase: PostgreSQL with Row Level Security, Auth, Storage, Edge Functions"],
        ["AI", "Google Gemini, called only through a server-side Edge Function"],
        ["Content", "22 bundled syllabus templates (JEE, NEET, CBSE, UPSC, GATE, CA, …)"],
        ["Size and tests", "About 36,000 lines of source; 2,283 automated tests"],
    ], widths=[1.4, 4.85], size=10)


    R.h1("The Problem")
    R.para(
        "Students track syllabi of hundreds of topics on paper, in spreadsheets or from memory, "
        "and switch between separate apps for timing, revision and AI help. Many study where "
        "connectivity is poor. Worse, the obvious way to compute progress — averaging chapter "
        "percentages — is wrong: a subject with a finished 1-topic chapter and an untouched "
        "99-topic chapter shows as 50% done when the truth is 1%. PrepPilot computes progress "
        "topic-weighted and shows an empty chapter as — rather than 0%."
    )

    # ---------------------------------------------------------- 2
    R.h1("Key Features")
    R.bullets([
        "**Syllabus tracker** with progress rings and bars at every level, several courses, "
        "22 templates and a Build-from-scratch form.",
        "**Timers:** stopwatch, Pomodoro (25/5/15) and custom, plus a full-screen landscape "
        "timer with six faces; every session is recorded against the syllabus.",
        "**Dashboard and sessions:** completion, today, total time, streak, last 7 days.",
        "**Revision:** flashcards (1d / 2d / 4d), quizzes of 1–15 questions at four "
        "difficulties saved to a Quiz deck, and any number of notes per topic.",
        "**AI assistant:** explain, summarise, quiz, answer questions about a photo, and "
        "suggest syllabus additions that the student reviews before they are added.",
        "**Focus Mode** (1–240 min) and student-created reminders.",
        "**Downloadable themes:** Classic built in; Notebook, Flight deck and Bold bento "
        "downloaded only when chosen.",
        "**Accounts and sync:** email, Google and guest sign-in; automatic cloud sync.",
    ])

    # ---------------------------------------------------------- 3
    R.h1("How It Works")
    R.para(
        "The SQLite database on the phone is the source of truth: every action is saved there "
        "first, so nothing waits for the network. A sync engine then pulls changes from "
        "Supabase, merges them by fixed rules — a deletion wins, a topic's tick is resolved on "
        "its own timestamp so a rename elsewhere cannot undo it, otherwise the later edit "
        "wins — and pushes the result. It runs on sign-in, every five minutes, on leaving the "
        "app and on reconnecting. Row Level Security in PostgreSQL keeps each student's data "
        "private. The Gemini key exists only on the server, and the assistant's suggestions "
        "can only add chapters or topics, never delete or complete anything."
    )
    R.figure(diagram("architecture.png"), "System architecture", width=Inches(3.6))

    # ---------------------------------------------------------- 4
    R.h1("Screens")
    R.para("Screens marked “captured 24 Sep 2026” are from the repository; those marked "
           "“web demo, 6 Oct 2026” were captured from the browser build, which uses sample "
           "data.", size=10, italic=True)
    R.grid([
        (shot("app/tracker.png"), "Tracker", REPO),
        (shot("app/subject-detail.png"), "Subject expanded into chapters", REPO),
        (shot("app/dashboard.png"), "Dashboard", REPO),
        (shot("app/templates.png"), "Template library", REPO),
        (shot("web/timer-running.png"), "Pomodoro timer", WEB),
        (shot("web/focus-setup.png"), "Focus Mode set-up", WEB),
        (shot("app/flashcard-review.png"), "Flashcard review", REPO),
        (shot("web/quiz-settings.png"), "Quiz settings", WEB),
        (shot("app/assistant-picture-reply.png"), "Assistant answering about a photo", REPO),
        (shot("app/assistant-proposal-review.png"), "Reviewing an AI suggestion", REPO),
        (shot("app/notes.png"), "Notes", REPO),
        (shot("web/settings-themes.png"), "Themes in Settings", WEB),
    ])
    R.grid([
        (shot("web/fs-paused-flip.png"), "Full-screen timer: Flip cards", WEB),
        (shot("web/fs-ring.png"), "Ring face", WEB),
        (shot("web/fs-hourglass.png"), "Hourglass face", WEB),
    ], cols=3, width=LAND_W)
    R.grid([
        (shot("web/theme-notebook-tracker.png"), "Notebook theme", WEB),
        (shot("web/theme-flightdeck-tracker.png"), "Flight deck theme", WEB),
        (shot("web/theme-bento-tracker.png"), "Bold bento theme", WEB),
        (shot("web/theme-notebook-light-tracker.png"), "Notebook, light mode", WEB),
    ])

    # ---------------------------------------------------------- 5
    R.h1("Testing and Results")
    R.table("Automated tests", [
        ["Package", "Tests", "Result"],
        ["App (Jest)", "1,658 in 86 suites", "All passed, run 6 Oct 2026"],
        ["Shared logic (Vitest)", "573 in 17 files", "All passed, run 6 Oct 2026"],
        ["Edge Functions (Deno)", "52 in 2 files", "Counted from source; run in CI"],
    ], widths=[1.9, 1.8, 2.55], size=10)
    R.para(
        "Continuous integration checks formatting and types, enforces coverage thresholds (95% "
        "statements and lines, 90% branches), and confirms the Android and web bundles build. "
        "Database tests run the app's real migrations against SQLite, and the sync engine is "
        "tested against an in-memory server that behaves like Supabase. According to the "
        "project documentation, the Android APK (about 41 MB) has been tested on a physical "
        "phone against the live backend, including Google sign-in, restoring data after a "
        "reinstall and photo questions. All four project objectives — accurate measurement, "
        "offline-first data, safe AI and one integrated workflow — are met."
    )

    # ---------------------------------------------------------- 6
    R.h1("Limitations and Next Steps")
    R.bullets([
        "Flashcards use fixed intervals; SM-2 statistics are kept so adaptive scheduling can be "
        "switched on later.",
        "Focus Mode cannot block other apps or turn on Do Not Disturb — Android does not allow "
        "it — and the app says so.",
        "The local database is not encrypted, and the AI rate limit is held in memory.",
        "Gemini runs on the free tier; a real launch needs the paid tier.",
        "Next: adaptive flashcards, local encryption, a durable rate limit, wider device "
        "testing and a Play Store release.",
    ])

    R.h1("Conclusion")
    R.para(
        "PrepPilot delivers a study tracker that works offline, measures progress honestly and "
        "adds AI help without giving up control of the student's work. It runs on Android and in "
        "the browser, and is backed by more than two thousand automated tests. The full "
        "Detailed Project Report covers the design, implementation and testing in depth."
    )
    return R


def main():
    if shutil.which("soffice") is None:
        sys.exit("soffice is required")
    R = build()
    R.save(OUT_DOCX)
    to_pdf(OUT_DOCX, DOCS)
    pages = subprocess.run(["pdfinfo", OUT_PDF], capture_output=True, text=True).stdout
    count = next((line.split()[-1] for line in pages.splitlines() if line.startswith("Pages")), "?")
    print(f"Wrote {OUT_DOCX}\nWrote {OUT_PDF}\n{count} pages")


if __name__ == "__main__":
    main()
