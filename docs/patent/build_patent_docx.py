"""
Renders PATENT_DRAFT.md to PrepPilot_Patent_Draft.docx.

The Markdown file is the single source: edit it, never the .docx, then run

    python docs/patent/build_patent_docx.py

It understands the subset of Markdown the draft uses — headings, paragraphs with
bold/italic/code, numbered and bulleted items (with indented continuation lines,
as in the claims), pipe tables, images with a width hint in the title
(`![caption](path "w=15")`, centimetres; two on one line sit side by side),
`<div align="center">` blocks and `<!-- pagebreak -->`.

Page setup follows common Indian patent-specification practice — A4, margins of
4 cm top and left and 3 cm bottom and right, Times New Roman 12 pt at 1.5 line
spacing, numbered pages. The agent should confirm it against the Patents Rules
in force when filing.
"""

import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE = os.path.join(HERE, "PATENT_DRAFT.md")
OUT = os.path.join(HERE, "PrepPilot_Patent_Draft.docx")

FONT = "Times New Roman"
BODY_PT = 12

# Prettier rewrites *italic* as _italic_ and "titles" as 'titles', so both are read.
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|(?<![\w\\])_[^_\s][^_]*_(?!\w)|`[^`]+`)")
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)(?:\s+[\"']w=([\d.]+)[\"'])?\)")


def set_font(run, size=BODY_PT, bold=None, italic=None, mono=False):
    run.font.name = "Consolas" if mono else FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), run.font.name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(par, text, size=BODY_PT, bold=False, italic=False):
    text = text.replace("&nbsp;", "\u00a0")
    # Prettier escapes Markdown punctuation (\_ for underscores); show the character.
    text = re.sub(r"\\([_*\[\]()#>|-])", r"\1", text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            set_font(r, size, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            set_font(r, size - 1, bold=bold, italic=italic, mono=True)
        elif part[0] in "*_" and part.endswith(part[0]) and len(part) > 2:
            r = par.add_run(part[1:-1])
            set_font(r, size, bold=bold, italic=True)
        else:
            r = par.add_run(part)
            set_font(r, size, bold=bold, italic=italic)


def spacing(par, before=0, after=6, line=1.5):
    fmt = par.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def page_number_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_font(run, 10)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("CONFIDENTIAL — DRAFT PATENT PROPOSAL — PrepPilot")
    set_font(hr, 9, italic=True)
    hr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def cell_border_none(table):
    tbl = table._tbl
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl.tblPr.append(borders)


def table_from_rows(doc, rows, centered):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1 + len(body), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate([header] + body):
        for c in range(len(header)):
            text = row[c] if c < len(row) else ""
            cell = table.cell(r, c)
            par = cell.paragraphs[0]
            spacing(par, 2, 2, 1.15)
            add_inline(par, text, size=10.5, bold=(r == 0))
            if r == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:val"), "clear")
                shade.set(qn("w:color"), "auto")
                shade.set(qn("w:fill"), "E7E6E6")
                cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph()


def image_path(rel):
    return os.path.normpath(os.path.join(HERE, rel))


def add_images(doc, images):
    """One image centred with its caption, or two side by side."""
    if len(images) == 1:
        caption, path, width = images[0]
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(par, 6, 2, 1.0)
        par.add_run().add_picture(image_path(path), width=Cm(float(width or 15)))
        if caption:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            spacing(cap, 2, 10, 1.0)
            add_inline(cap, caption, size=10.5, bold=True)
        return
    table = doc.add_table(rows=2, cols=len(images))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_border_none(table)
    for i, (caption, path, width) in enumerate(images):
        pic = table.cell(0, i).paragraphs[0]
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(image_path(path), width=Cm(float(width or 6)))
        cap = table.cell(1, i).paragraphs[0]
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(cap, 2, 8, 1.0)
        add_inline(cap, caption, size=9.5, bold=True)


def build():
    lines = open(SOURCE, encoding="utf-8").read().splitlines()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.top_margin, section.left_margin = Cm(4), Cm(4)
    section.bottom_margin, section.right_margin = Cm(3), Cm(3)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    page_number_footer(section)

    centered = False
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith('<div align="center">'):
            centered = True
            i += 1
            continue
        if stripped == "</div>":
            centered = False
            i += 1
            continue
        if stripped == "<!-- pagebreak -->":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        if not stripped:
            i += 1
            continue

        # Images: one or more on a line.
        images = IMAGE.findall(stripped)
        if images and IMAGE.sub("", stripped).strip() == "":
            add_images(doc, images)
            i += 1
            continue

        # Tables.
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            table_from_rows(doc, rows, centered)
            continue

        # Headings.
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            par = doc.add_paragraph()
            if centered:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                size = {1: 22, 2: 15, 3: 13}[level]
                spacing(par, 10, 10, 1.15)
            else:
                size = {1: 16, 2: 14, 3: 12.5}[level]
                spacing(par, 12 if level > 1 else 0, 8, 1.15)
                par.paragraph_format.keep_with_next = True
                # Built-in heading outline so Word's navigation pane lists them.
                ppr = par._p.get_or_add_pPr()
                outline = OxmlElement("w:outlineLvl")
                outline.set(qn("w:val"), str(level - 1))
                ppr.append(outline)
            add_inline(par, text.replace("**", ""), size=size, bold=True)
            i += 1
            continue

        # Numbered items (claims, declarations, merge rules) with continuations.
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m and not raw.startswith(" "):
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt = par.paragraph_format
            fmt.left_indent = Cm(1.0)
            fmt.first_line_indent = Cm(-1.0)
            spacing(par, 0, 4, 1.5)
            add_inline(par, f"{m.group(1)}.\t{m.group(2)}")
            i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                sub.paragraph_format.left_indent = Cm(1.8)
                sub.paragraph_format.first_line_indent = Cm(-0.8)
                spacing(sub, 0, 3, 1.5)
                add_inline(sub, lines[i].strip())
                i += 1
            continue

        # Bullets.
        if stripped.startswith("- "):
            par = doc.add_paragraph()
            fmt = par.paragraph_format
            fmt.left_indent = Cm(0.9)
            fmt.first_line_indent = Cm(-0.5)
            spacing(par, 0, 3, 1.5)
            add_inline(par, "•  " + stripped[2:])
            i += 1
            continue

        # Ordinary paragraph.
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.JUSTIFY
        spacing(par, 0, 8 if not centered else 6, 1.15 if centered else 1.5)
        add_inline(par, stripped, size=BODY_PT + (1 if centered else 0))
        i += 1

    doc.core_properties.title = "Draft Proposal for Patent Filing — PrepPilot"
    doc.core_properties.author = "Archit Jain; Vanshika Rajput; Lavi Chaudhary; Jasmine Chaudhary"
    doc.save(OUT)
    print("Saved:", OUT)


if __name__ == "__main__":
    build()
